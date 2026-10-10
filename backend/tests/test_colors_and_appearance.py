"""Tests for cell fill (background) and sheet appearance (grid color, banding)."""

import os
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from backend.controller.controller import Controller
from backend.document.document import Document
from backend.models.cell_style import NO_FILL, CellStyle
from backend.models.sheet_appearance import SheetAppearance
from backend.network.client import NetworkClient
from backend.storage.excel_storage import ExcelStorage
from backend.tests.test_text_styles import make_controller, reopen, styles
from backend.utils.config import SERVER_IP, SERVER_PORT

APPEARANCE_NAME = "_excelence_appearance"


def fill_of(path, cell, sheet="Sheet1"):
    fill = load_workbook(path)[sheet][cell].fill
    return fill.fill_type, (fill.start_color.rgb[-6:] if fill.fill_type else None)


# ----------------------------------------------------------------------
# Value objects
# ----------------------------------------------------------------------

def test_background_validation():
    empty = CellStyle()
    assert empty.apply_patch({"background": "#ffcc00"}).background == "FFCC00"
    assert empty.apply_patch({"background": "none"}).background == NO_FILL
    assert empty.apply_patch({"background": "NONE"}).background == NO_FILL

    for bad in ("red", "12345", "GGGGGG", 123456, True, "", "#"):
        assert empty.apply_patch({"background": bad}) is None, bad

    # reset returns it to inherit
    assert CellStyle(background="FFCC00").apply_patch({}, ["background"]).background is None


def test_background_merges_like_other_fields():
    merged = CellStyle(background="111111").merged_over(CellStyle(background="222222", bold=True))
    assert (merged.background, merged.bold) == ("111111", True)
    assert CellStyle().merged_over(CellStyle(background="222222")).background == "222222"
    # an explicit "no fill" override beats a default fill
    assert CellStyle(background=NO_FILL).merged_over(CellStyle(background="222222")).background == NO_FILL


def test_appearance_validation_and_patch():
    empty = SheetAppearance()
    patched = empty.apply_patch({"grid_color": "#c0c0c0", "banding": True, "banding_color": "f2f2f2"})
    assert patched.to_dict() == {"grid_color": "C0C0C0", "banding": True, "banding_color": "F2F2F2"}

    for bad in (
        {"grid_color": "red"}, {"banding": "yes"}, {"banding": 1}, {"banding_color": 5},
        {"nope": True}, {"grid_color": "NONE"},
    ):
        assert empty.apply_patch(bad) is None, bad

    assert patched.apply_patch({}, ["grid_color"]).grid_color is None
    assert empty.apply_patch({"banding": True}, ["banding"]) is None      # set and reset together
    assert SheetAppearance.from_dict({"grid_color": "bad", "banding": True}).to_dict() == {"banding": True}


# ----------------------------------------------------------------------
# Cell fill through the controller and the file
# ----------------------------------------------------------------------

def test_cell_fill_is_a_real_excel_fill_and_round_trips():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_cell_style(0, 0, 1, 1, {"background": "FFF2CC"}) is True
        assert c.set_cell_style(1, 1, 1, 1, {"bold": True, "color": "C00000"}) is True

        path = os.path.join(folder, "t.xlsx")
        fresh = reopen(c, folder)

        assert fill_of(path, "A2") == ("solid", "FFF2CC")
        assert fill_of(path, "B3") == ("solid", "FFF2CC")
        assert fill_of(path, "C2") == (None, None)                 # untouched cell
        assert fill_of(path, "A1") == (None, None)                 # header row untouched

        assert styles(fresh)[(0, 0)] == {"background": "FFF2CC"}
        assert styles(fresh)[(1, 1)] == {"background": "FFF2CC", "bold": True, "color": "C00000"}
        assert len(styles(fresh)) == 4


def test_clearing_a_fill_removes_it_from_the_file():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(0, 0, 0, 0, {"background": "FFF2CC"})
        c.set_cell_style(0, 0, 0, 0, {}, ["background"])
        assert c.save_document() is True
        assert fill_of(os.path.join(folder, "t.xlsx"), "A2") == (None, None)


def test_default_fill_applies_everywhere_and_cells_can_opt_out_or_override():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_text_defaults({"background": "DDEBF7"}) is True
        c.set_cell_style(0, 0, 0, 0, {"background": "FFF2CC"})     # override
        c.set_cell_style(1, 0, 1, 0, {"background": NO_FILL})      # explicit no fill

        path = os.path.join(folder, "t.xlsx")
        fresh = reopen(c, folder)

        assert fill_of(path, "B2") == ("solid", "DDEBF7")          # default
        assert fill_of(path, "A2") == ("solid", "FFF2CC")          # override wins
        assert fill_of(path, "A3") == (None, None)                 # opted out
        assert fill_of(path, "A1") == (None, None)                 # header untouched

        assert fresh.text_defaults() == {"background": "DDEBF7"}
        assert styles(fresh) == {
            (0, 0): {"background": "FFF2CC"},
            (1, 0): {"background": NO_FILL},
        }


def test_no_fill_is_rejected_as_a_sheet_default():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_text_defaults({"background": NO_FILL}) is False
        assert c.text_defaults() == {}


def test_fill_is_part_of_undo():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(0, 0, 0, 0, {"background": "FFF2CC"})
        assert c.undo() is True
        assert styles(c) == {}
        fresh = reopen(c, folder)
        assert styles(fresh) == {}
        assert fill_of(os.path.join(folder, "t.xlsx"), "A2") == (None, None)


def test_fills_follow_rows_through_structural_edits():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(1, 0, 1, 0, {"background": "E2EFDA"})     # Bob
        c.insert_row(0)
        assert styles(c) == {(2, 0): {"background": "E2EFDA"}}
        fresh = reopen(c, folder)
        assert styles(fresh) == {(2, 0): {"background": "E2EFDA"}}


def test_fills_from_a_workbook_made_elsewhere_are_imported():
    with tempfile.TemporaryDirectory() as folder:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(["Name", "Age"])
        sheet.append(["Alice", 25])
        sheet.append(["Bob", 30])
        sheet["A2"].fill = PatternFill(fill_type="solid", start_color="FFFFFF00", end_color="FFFFFF00")
        sheet["B2"].font = Font(bold=True)
        sheet["B3"].fill = PatternFill(fill_type="gray125")        # pattern fill: ignored
        workbook.save(os.path.join(folder, "ext.xlsx"))

        document = Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT))
        document.documents_folder = folder
        c = Controller(document)
        assert c.open_document("ext.xlsx") is True
        assert styles(c) == {(0, 0): {"background": "FFFF00"}, (0, 1): {"bold": True}}


# ----------------------------------------------------------------------
# Appearance
# ----------------------------------------------------------------------

def test_appearance_set_reset_and_validation():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.appearance() == {}

        assert c.set_appearance({"grid_color": "#C0C0C0", "banding": True, "banding_color": "F2F2F2"}) is True
        assert c.appearance() == {"grid_color": "C0C0C0", "banding": True, "banding_color": "F2F2F2"}

        assert c.set_appearance({}, ["grid_color"]) is True
        assert c.appearance() == {"banding": True, "banding_color": "F2F2F2"}

        before = c.appearance()
        assert c.set_appearance({"grid_color": "nope"}) is False
        assert c.set_appearance({"banding": "yes"}) is False
        assert c.set_appearance({}, ["bogus"]) is False
        assert c.appearance() == before


def test_appearance_persists_and_is_per_sheet():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_appearance({"grid_color": "AABBCC", "banding": True})

        assert c.add_sheet("Other") is True
        assert c.set_sheet("Other") is True
        assert c.appearance() == {}                                # not shared
        c.set_appearance({"banding_color": "112233"})
        assert c.set_sheet("Sheet1") is True
        assert c.appearance() == {"grid_color": "AABBCC", "banding": True}

        assert c.rename_sheet("Sheet1", "Renamed") is True
        assert c.appearance() == {"grid_color": "AABBCC", "banding": True}

        fresh = reopen(c, folder)
        fresh.set_sheet("Renamed")
        assert fresh.appearance() == {"grid_color": "AABBCC", "banding": True}
        fresh.set_sheet("Other")
        assert fresh.appearance() == {"banding_color": "112233"}


def test_clearing_appearance_removes_the_marker_from_the_file():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_appearance({"banding": True})
        c.set_appearance({}, ["banding"])
        assert c.save_document() is True
        assert APPEARANCE_NAME not in load_workbook(os.path.join(folder, "t.xlsx"))["Sheet1"].defined_names


def test_appearance_is_not_undoable_and_survives_undo():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.edit_cell(0, 0, "x")
        c.set_appearance({"banding": True})
        assert len(c.document.undo_stack) == 1                     # the edit only
        assert c.undo() is True
        assert c.appearance() == {"banding": True}


def test_opening_another_document_does_not_inherit_appearance():
    with tempfile.TemporaryDirectory() as folder:
        a = make_controller(folder, "a.xlsx")
        a.set_appearance({"banding": True, "grid_color": "AABBCC"})
        plain = Workbook()
        plain.active.title = "Sheet1"
        plain.active.append(["x"])
        plain.active.append([1])
        plain.save(os.path.join(folder, "b.xlsx"))
        assert a.open_document("b.xlsx") is True
        assert a.appearance() == {}


def test_corrupt_appearance_marker_does_not_block_opening():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_appearance({"banding": True})
        assert c.save_document() is True

        path = os.path.join(folder, "t.xlsx")
        workbook = load_workbook(path)
        workbook["Sheet1"].defined_names[APPEARANCE_NAME].attr_text = '"{broken'
        workbook.save(path)

        fresh = Controller(Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT)))
        fresh.document.documents_folder = folder
        assert fresh.open_document("t.xlsx") is True
        assert fresh.appearance() == {}


def test_grid_payload_carries_appearance_and_fills():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_appearance({"banding": True})
        c.set_cell_style(0, 0, 0, 0, {"background": "FFF2CC"})
        data = c.grid_data()
        assert data.appearance == {"banding": True}
        assert data.cell_styles == [{"row": 0, "column": 0, "style": {"background": "FFF2CC"}}]
