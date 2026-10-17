"""What a packaged (PyInstaller) server depends on, tested without PyInstaller."""

import os
import subprocess
import sys
import tempfile
import textwrap

from openpyxl import Workbook, load_workbook

from backend.run_server import STARTER_DOCUMENT, ensure_starter_document

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def test_a_fresh_install_gets_a_loadable_starter_document():
    with tempfile.TemporaryDirectory() as parent:
        folder = os.path.join(parent, "does", "not", "exist", "yet")
        assert ensure_starter_document(folder) is True
        path = os.path.join(folder, STARTER_DOCUMENT)
        assert os.path.isfile(path)
        assert load_workbook(path).sheetnames                      # a real workbook

        from backend.tests.test_text_styles import Controller, Document, ExcelStorage, NetworkClient
        from backend.utils.config import SERVER_IP, SERVER_PORT
        document = Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT))
        document.documents_folder = folder
        assert Controller(document).open_document(STARTER_DOCUMENT) is True


def test_an_existing_documents_folder_is_never_touched():
    with tempfile.TemporaryDirectory() as folder:
        book = Workbook()
        book.save(os.path.join(folder, "mine.xlsx"))
        before = sorted(os.listdir(folder))
        assert ensure_starter_document(folder) is False
        assert sorted(os.listdir(folder)) == before


def test_a_workbook_in_a_subfolder_also_counts():
    with tempfile.TemporaryDirectory() as folder:
        os.makedirs(os.path.join(folder, "reports"))
        Workbook().save(os.path.join(folder, "reports", "q1.xlsx"))
        assert ensure_starter_document(folder) is False
        assert not os.path.exists(os.path.join(folder, STARTER_DOCUMENT))


def test_other_files_do_not_count_as_documents():
    with tempfile.TemporaryDirectory() as folder:
        open(os.path.join(folder, "notes.txt"), "w").close()
        open(os.path.join(folder, ".excelence_state.json"), "w").write("{}")
        assert ensure_starter_document(folder) is True


def _run_as_frozen(executable_dir: str, extra_env: dict | None = None) -> dict:
    """Import the server's config/logger in a process that claims to be frozen."""
    script = textwrap.dedent(f"""
        import json, os, sys
        sys.frozen = True
        sys.executable = {os.path.join(executable_dir, 'excelence-server.exe')!r}
        sys.path.insert(0, {PROJECT_ROOT!r})
        import backend.utils.config as config
        import backend.utils.logger as logger_module
        logger_module.get_logger("probe_frozen").info("hello")
        print(json.dumps({{
            "base": config.BASE_DIR,
            "docs": config.DOCUMENTS_FOLDER,
            "log_exists": os.path.isfile(os.path.join({executable_dir!r}, "logs", "backend.log")),
        }}))
    """)
    env = dict(os.environ)
    env.pop("EXCELENCE_DOCUMENTS_DIR", None)
    env.update(extra_env or {})
    result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, env=env, cwd=executable_dir)
    assert result.returncode == 0, result.stderr
    import json
    return json.loads(result.stdout.strip().splitlines()[-1])


def test_a_frozen_server_keeps_data_and_logs_next_to_the_executable():
    with tempfile.TemporaryDirectory() as folder:
        info = _run_as_frozen(folder)
        assert info["base"] == folder
        assert info["docs"] == os.path.join(folder, "data", "documents")   # survives restarts, unlike a temp dir
        assert info["log_exists"] is True                                  # not buried inside the bundle


def test_the_documents_folder_can_still_be_moved_when_frozen():
    with tempfile.TemporaryDirectory() as folder, tempfile.TemporaryDirectory() as elsewhere:
        info = _run_as_frozen(folder, {"EXCELENCE_DOCUMENTS_DIR": elsewhere})
        assert info["docs"] == elsewhere
