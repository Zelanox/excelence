"""Tests for sheet-default and per-cell text styles.

Written as plain functions (no fixtures) like the other tests here, so they
run under pytest or a bare loop. Each test builds its own workbook in a
temp folder and points a Document at it.
"""

import json
import os
import tempfile
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from backend.controller.controller import Controller
from backend.document.document import Document
from backend.models.cell_style import CellStyle
from backend.network.client import NetworkClient
from backend.storage.excel_storage import ExcelStorage
from backend.utils.config import SERVER_IP, SERVER_PORT

DEFAULTS_NAME = "_excelence_text_defaults"

ROWS = [
    ["Alice", 25, 90],
    ["Bob", 30, 85],
    ["Charlie", 22, 95],
    ["Dana", 41, 70],
]


def make_controller(folder: str, name: str = "t.xlsx", rows=ROWS, header=("Name", "Age", "Score")) -> Controller:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(list(header))

    for row in rows:
        sheet.append(list(row))

    workbook.save(os.path.join(folder, name))

    document = Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT))
    document.documents_folder = folder
    controller = Controller(document)
    assert controller.open_document(name) is True
    return controller


def styles(controller: Controller) -> dict:
    return {
        (entry["row"], entry["column"]): entry["style"]
        for entry in controller.cell_styles()
    }


def reopen(controller: Controller, folder: str, name: str = "t.xlsx") -> Controller:
    assert controller.save_document() is True
    document = Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT))
    document.documents_folder = folder
    fresh = Controller(document)
    assert fresh.open_document(name) is True
    return fresh


# ----------------------------------------------------------------------
# CellStyle value object
# ----------------------------------------------------------------------

def test_cell_style_patch_and_merge():
    base = CellStyle(bold=True, font_family="Cairo", font_size=12.0)
    patched = base.apply_patch({"italic": True, "color": "#ff0000"}, ["bold"])

    assert patched.to_dict() == {
        "italic": True,
        "font_family": "Cairo",
        "font_size": 12.0,
        "color": "FF0000",
    }
    # original untouched
    assert base.bold is True

    merged = CellStyle(italic=True).merged_over(base)
    assert merged.bold is True and merged.italic is True and merged.font_family == "Cairo"


def test_cell_style_rejects_invalid_values():
    empty = CellStyle()

    for bad in (
        {"bold": "yes"},
        {"bold": 1},
        {"font_size": 0},
        {"font_size": 0.5},
        {"font_size": 410},
        {"font_size": True},
        {"font_size": float("nan")},
        {"font_size": "12"},
        {"font_family": ""},
        {"font_family": "   "},
        {"font_family": "x" * 65},
        {"font_family": "bad\nname"},
        {"font_family": 5},
        {"color": "red"},
        {"color": "#12345"},
        {"color": "GGGGGG"},
        {"color": 123456},
        {"nope": True},
    ):
        assert empty.apply_patch(bad) is None, bad

    assert empty.apply_patch({}, ["nope"]) is None
    # a field cannot be both set and reset
    assert empty.apply_patch({"bold": True}, ["bold"]) is None
    # boundaries are valid
    assert empty.apply_patch({"font_size": 1}) is not None
    assert empty.apply_patch({"font_size": 409}) is not None


def test_from_dict_is_lenient():
    style = CellStyle.from_dict({"bold": True, "font_size": "junk", "color": "zz", "extra": 1})
    assert style.to_dict() == {"bold": True}
    assert CellStyle.from_dict("not a dict").is_empty()
    assert CellStyle.from_dict(None).is_empty()


# ----------------------------------------------------------------------
# Per-cell styles through the controller
# ----------------------------------------------------------------------

def test_set_cell_style_range_and_data_payload():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)

        # corners given in reverse order must still work
        assert c.set_cell_style(2, 1, 0, 0, {"bold": True, "color": "00AA00"}) is True

        found = styles(c)
        assert set(found) == {(r, col) for r in range(3) for col in range(2)}
        assert all(s == {"bold": True, "color": "00AA00"} for s in found.values())
        assert c.text_defaults() == {}


def test_reset_and_empty_cleanup():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_cell_style(0, 0, 0, 0, {"bold": True, "italic": True}) is True
        assert c.set_cell_style(0, 0, 0, 0, {}, ["bold"]) is True
        assert styles(c) == {(0, 0): {"italic": True}}

        assert c.set_cell_style(0, 0, 0, 0, {}, ["italic"]) is True
        assert styles(c) == {}
        assert c.document.spreadsheet.current_sheet().cell_styles == {}


def test_invalid_requests_change_nothing():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_cell_style(0, 0, 0, 0, {"bold": True}) is True
        before = styles(c)

        assert c.set_cell_style(0, 0, 1, 1, {"color": "nope"}) is False
        assert c.set_cell_style(-1, 0, 0, 0, {"bold": True}) is False
        assert c.set_cell_style(0, 0, 4, 0, {"bold": True}) is False    # row out of range
        assert c.set_cell_style(0, 0, 0, 3, {"bold": True}) is False    # column out of range
        assert c.set_cell_style(0, 0, 1, 1, {"italic": True}, ["italic"]) is False
        assert c.set_cell_style(True, 0, 0, 0, {"bold": True}) is False
        assert c.set_cell_style(0, 0, 0, 0, None, ["nope"]) is False

        assert styles(c) == before


def test_style_survives_edit_save_and_reload_as_real_fonts():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_cell_style(
            1, 0, 1, 0,
            {"bold": True, "italic": True, "underline": True, "strikethrough": True,
             "font_family": "Cairo", "font_size": 16, "color": "FF0000"},
        ) is True

        # an unrelated edit rebuilds the whole worksheet - style must survive
        assert c.edit_cell(2, 1, 99) is True
        assert styles(c)[(1, 0)]["font_family"] == "Cairo"

        fresh = reopen(c, folder)
        assert styles(fresh) == {
            (1, 0): {
                "bold": True, "italic": True, "underline": True, "strikethrough": True,
                "font_family": "Cairo", "font_size": 16.0, "color": "FF0000",
            }
        }

        # and it is a REAL Excel font, in the right cell (Bob = excel row 3, col A)
        font = load_workbook(os.path.join(folder, "t.xlsx"))["Sheet1"]["A3"].font
        assert (font.name, font.sz, font.b, font.i, font.strike) == ("Cairo", 16.0, True, True, True)
        assert font.u == "single"
        assert font.color.rgb[-6:] == "FF0000"


def test_removed_override_leaves_no_stale_font_in_file():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(0, 0, 0, 0, {"bold": True, "font_family": "Cairo"})
        c.set_cell_style(0, 0, 0, 0, {}, ["bold", "font_family"])
        assert c.save_document() is True

        font = load_workbook(os.path.join(folder, "t.xlsx"))["Sheet1"]["A2"].font
        assert font.name == "Calibri" and not font.b


# ----------------------------------------------------------------------
# Sheet defaults
# ----------------------------------------------------------------------

def test_text_defaults_roundtrip_and_override_precedence():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_text_defaults({"font_family": "Amiri", "font_size": 14, "italic": True}) is True
        assert c.set_cell_style(0, 0, 0, 0, {"bold": True, "font_size": 20}) is True

        assert c.text_defaults() == {"italic": True, "font_family": "Amiri", "font_size": 14.0}

        fresh = reopen(c, folder)
        assert fresh.text_defaults() == {"italic": True, "font_family": "Amiri", "font_size": 14.0}
        # only the DIFFERENCES from the default come back as overrides
        assert styles(fresh) == {(0, 0): {"bold": True, "font_size": 20.0}}

        sheet = load_workbook(os.path.join(folder, "t.xlsx"))["Sheet1"]
        # default applied to ordinary data cells so Excel looks right too...
        assert (sheet["B3"].font.name, sheet["B3"].font.sz, sheet["B3"].font.i) == ("Amiri", 14.0, True)
        # ...merged under the override
        assert (sheet["A2"].font.name, sheet["A2"].font.sz, sheet["A2"].font.b) == ("Amiri", 20.0, True)
        # header row left alone
        assert sheet["A1"].font.name == "Calibri"


def test_clearing_defaults_removes_marker_and_fonts():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_text_defaults({"font_family": "Amiri"})
        c.set_text_defaults({}, ["font_family"])
        assert c.text_defaults() == {}
        assert c.save_document() is True

        sheet = load_workbook(os.path.join(folder, "t.xlsx"))["Sheet1"]
        assert DEFAULTS_NAME not in sheet.defined_names
        assert sheet["B3"].font.name == "Calibri"


def test_invalid_defaults_rejected():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        assert c.set_text_defaults({"font_size": -3}) is False
        assert c.set_text_defaults({"bold": "x"}) is False
        assert c.set_text_defaults({}, ["bogus"]) is False
        assert c.text_defaults() == {}


def test_defaults_do_not_become_overrides_after_reload():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_text_defaults({"bold": True, "color": "112233"})
        fresh = reopen(c, folder)
        assert styles(fresh) == {}
        assert fresh.text_defaults() == {"bold": True, "color": "112233"}


# ----------------------------------------------------------------------
# Styles must follow the data through structural edits
# ----------------------------------------------------------------------

def test_insert_and_delete_row_shift_styles():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(1, 0, 1, 0, {"bold": True})     # Bob

        assert c.insert_row(0) is True                   # Bob is now row 2
        assert styles(c) == {(2, 0): {"bold": True}}
        assert c.cell_value(2, 0) == "Bob"

        assert c.insert_row(3) is True                   # below Bob: no shift
        assert styles(c) == {(2, 0): {"bold": True}}

        assert c.delete_row(0) is True                   # back to row 1
        assert styles(c) == {(1, 0): {"bold": True}}
        assert c.cell_value(1, 0) == "Bob"

        assert c.delete_row(1) is True                   # delete Bob: style gone
        assert styles(c) == {}
        assert c.document.spreadsheet.current_sheet().cell_styles == {}


def test_structural_edits_survive_save_and_reload():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(2, 0, 2, 0, {"italic": True})   # Charlie
        c.insert_row(0)
        fresh = reopen(c, folder)
        assert fresh.cell_value(3, 0) == "Charlie"
        assert styles(fresh) == {(3, 0): {"italic": True}}


def test_delete_and_rename_column_move_styles():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(0, 1, 0, 2, {"bold": True})     # Age + Score on row 0

        assert c.rename_column("Age", "Years") is True
        assert set(styles(c)) == {(0, 1), (0, 2)}

        assert c.delete_column("Name") is True           # Years/Score shift left
        assert set(styles(c)) == {(0, 0), (0, 1)}

        assert c.delete_column("Years") is True
        assert set(styles(c)) == {(0, 0)}

        # the last remaining column can't be deleted; nothing may change
        assert c.delete_column("Score") is False
        assert set(styles(c)) == {(0, 0)}


def test_insert_column_does_not_disturb_styles():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(0, 1, 0, 1, {"bold": True})     # Age
        assert c.insert_column("New", 0) is True
        assert styles(c) == {(0, 2): {"bold": True}}      # Age moved right by one


# ----------------------------------------------------------------------
# View coordinates: sort / filter
# ----------------------------------------------------------------------

def test_styles_follow_rows_through_sort_and_filter():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_cell_style(1, 0, 1, 0, {"bold": True})     # Bob (score 85)

        assert c.sort([{"column": "Score", "ascending": False}]) is True
        # sorted by score desc: Charlie 95, Alice 90, Bob 85, Dana 70
        assert c.cell_value(2, 0) == "Bob"
        assert styles(c) == {(2, 0): {"bold": True}}

        # formatting a VIEW row must hit the right underlying row
        assert c.set_cell_style(0, 0, 0, 0, {"italic": True}) is True   # Charlie
        assert c.clear_sort() is True
        assert c.cell_value(2, 0) == "Charlie"
        assert styles(c) == {(1, 0): {"bold": True}, (2, 0): {"italic": True}}

        assert c.search("Dana") is True
        assert c.filtered_row_count() == 1
        assert c.set_cell_style(0, 2, 0, 2, {"bold": True}) is True      # Dana's score
        assert styles(c) == {(0, 2): {"bold": True}}
        assert c.clear_search() is True
        assert styles(c)[(3, 2)] == {"bold": True}


# ----------------------------------------------------------------------
# Sheets, documents, foreign files
# ----------------------------------------------------------------------

def test_styles_are_per_sheet_and_survive_sheet_ops():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_text_defaults({"font_family": "Cairo"})
        c.set_cell_style(0, 0, 0, 0, {"bold": True})

        assert c.add_sheet("Other") is True
        assert c.set_sheet("Other") is True
        assert styles(c) == {} and c.text_defaults() == {}

        c.set_text_defaults({"italic": True})
        assert c.set_sheet("Sheet1") is True
        assert c.text_defaults() == {"font_family": "Cairo"}
        assert styles(c) == {(0, 0): {"bold": True}}

        assert c.rename_sheet("Sheet1", "Renamed") is True
        assert c.text_defaults() == {"font_family": "Cairo"}
        assert styles(c) == {(0, 0): {"bold": True}}

        fresh = reopen(c, folder)
        assert set(fresh.sheets()) == {"Renamed", "Other"}
        fresh.set_sheet("Other")
        assert fresh.text_defaults() == {"italic": True}
        fresh.set_sheet("Renamed")
        assert fresh.text_defaults() == {"font_family": "Cairo"}
        assert styles(fresh) == {(0, 0): {"bold": True}}


def test_opening_another_document_does_not_inherit_styles():
    with tempfile.TemporaryDirectory() as folder:
        a = make_controller(folder, "a.xlsx")
        a.set_text_defaults({"font_family": "Cairo"})
        a.set_cell_style(0, 0, 0, 0, {"bold": True})

        # same Document object opens a second, plain workbook that also
        # has a sheet named "Sheet1"
        other = Workbook()
        other.active.title = "Sheet1"
        other.active.append(["x", "y"])
        other.active.append([1, 2])
        other.save(os.path.join(folder, "b.xlsx"))

        assert a.open_document("b.xlsx") is True
        assert a.text_defaults() == {}
        assert styles(a) == {}


def test_styling_from_a_workbook_made_elsewhere_is_imported():
    with tempfile.TemporaryDirectory() as folder:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(["Name", "Age"])
        sheet.append(["Alice", 25])
        sheet.append(["Bob", 30])
        sheet["A2"].font = Font(name="Arial", sz=10, b=True, color="FF336699")
        sheet["B3"].font = Font(name="Calibri", sz=11)          # == baseline, no override
        workbook.save(os.path.join(folder, "ext.xlsx"))

        document = Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT))
        document.documents_folder = folder
        c = Controller(document)
        assert c.open_document("ext.xlsx") is True

        assert styles(c) == {(0, 0): {
            "bold": True, "font_family": "Arial", "font_size": 10.0, "color": "336699",
        }}
        assert c.text_defaults() == {}


def test_corrupt_defaults_marker_does_not_block_opening():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_text_defaults({"bold": True})
        assert c.save_document() is True

        path = os.path.join(folder, "t.xlsx")
        workbook = load_workbook(path)
        workbook["Sheet1"].defined_names[DEFAULTS_NAME].attr_text = '"{not json"'
        workbook.save(path)

        fresh = Controller(Document(ExcelStorage(), NetworkClient(SERVER_IP, SERVER_PORT)))
        fresh.document.documents_folder = folder
        assert fresh.open_document("t.xlsx") is True
        assert fresh.text_defaults() == {}


def test_grid_payload_is_one_consistent_snapshot():
    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder)
        c.set_rtl(True)
        c.set_column_width("Name", 30)
        c.set_text_defaults({"bold": True})
        c.set_cell_style(0, 0, 0, 0, {"italic": True})

        data = c.grid_data()
        assert data.rtl is True
        assert data.column_widths == {"Name": 30.0}
        assert data.text_defaults == {"bold": True}
        assert data.cell_styles == [{"row": 0, "column": 0, "style": {"italic": True}}]
        assert data.row_count == len(data.rows) == 4


# ----------------------------------------------------------------------
# Scale
# ----------------------------------------------------------------------

def test_large_sheet_timing_smoke():
    """Not a pass/fail perf gate - prints timings so growth is visible."""
    rows = [[f"name{i}", i, i * 2, i % 7, "x"] for i in range(20000)]

    with tempfile.TemporaryDirectory() as folder:
        c = make_controller(folder, rows=rows, header=("a", "b", "c", "d", "e"))

        start = time.perf_counter()
        assert c.set_text_defaults({"font_family": "Cairo", "font_size": 12}) is True
        defaults_seconds = time.perf_counter() - start

        start = time.perf_counter()
        assert c.set_cell_style(0, 0, 19999, 4, {"bold": True}) is True
        range_seconds = time.perf_counter() - start

        start = time.perf_counter()
        fresh = reopen(c, folder)
        reopen_seconds = time.perf_counter() - start

        assert len(styles(fresh)) == 100000
        print(
            f"20k rows x 5 cols: defaults {defaults_seconds:.2f}s, "
            f"whole-sheet bold {range_seconds:.2f}s, save+reopen {reopen_seconds:.2f}s"
        )
